import docx
import shutil

src_template = "/Users/ProM2/Documents/Lab_Pareto_30_phut_De_sinh_vien.docx"
target_pareto = "/Users/ProM2/Documents/Pareto.docx"
target_artifact = "/Users/ProM2/.gemini/antigravity/brain/23722a61-f5ab-4410-9ac7-d23db12af167/.user_uploaded/media_1790672577410.docx"

doc = docx.Document(src_template)

# Header block
doc.paragraphs[0].text = "Bài thực hành Pareto trong quản lý chất lượng dự án phần mềm SmartRecruit"
doc.paragraphs[0].runs[0].font.name = "Arial"
doc.paragraphs[0].runs[0].font.size = docx.shared.Pt(18)
doc.paragraphs[0].runs[0].font.bold = True
doc.paragraphs[0].runs[0].font.color.rgb = docx.shared.RGBColor(0x1F, 0x4E, 0x79)

doc.paragraphs[1].text = "Học phần: Quản lý dự án CNTT  |  Thời gian: 30 phút  |  Đơn vị thực hiện: [Nhóm 03] (5 sinh viên)"
doc.paragraphs[1].runs[0].font.name = "Arial"
doc.paragraphs[1].runs[0].font.size = docx.shared.Pt(10)
doc.paragraphs[1].runs[0].font.italic = True
doc.paragraphs[1].runs[0].font.color.rgb = docx.shared.RGBColor(0x59, 0x59, 0x59)

doc.paragraphs[2].text = "Lớp: Chuyên ngành Kỹ thuật Phần mềm – Khóa 2023  |  Nhóm: [Nhóm 03]  |  Ngày: 05/10/2026 (Cuối Sprint 4)"
doc.paragraphs[2].runs[0].font.name = "Arial"
doc.paragraphs[2].runs[0].font.size = docx.shared.Pt(10)
doc.paragraphs[2].runs[0].font.bold = True
doc.paragraphs[2].runs[0].font.color.rgb = docx.shared.RGBColor(0x1F, 0x4E, 0x79)

# 1 Tình huống và nhiệm vụ
doc.paragraphs[4].text = "Nhóm thực hiện Web MVP Nền tảng tuyển dụng thông minh SmartRecruit trong 8 tuần (từ 08/09/2026 đến 02/11/2026), mỗi tuần một Sprint. Cuối Sprint 4, đội ngũ QA và phát triển ghi nhận 100 báo cáo sự cố (Defect Reports) thực tế từ các phân hệ trích xuất CV qua Gemini 2.5 Flash, xác thực phân quyền React 19 / JWT, tải file S3 và tính điểm khớp kỹ năng. Hãy phân tích Pareto trên Excel và đề xuất 3 công việc cải thiện chất lượng ưu tiên cho Sprint 5."
doc.paragraphs[5].text = "Mục tiêu: Tính tỷ lệ phần trăm và tỷ lệ lũy kế, vẽ biểu đồ Pareto kết hợp (Combo Chart), nhận diện nhóm lỗi trọng yếu (Vital Few) chiếm 80% sự cố hệ thống và giải thích quyết định ưu tiên kết hợp đánh giá mức độ nghiêm trọng bảo mật (Severity). Trong 30 phút nhóm phân tích số liệu và lập kế hoạch hành động Sprint 5 có tiêu chí nghiệm thu kiểm chứng được."

# 2 Dữ liệu thực hành
doc.paragraphs[7].text = "Dữ liệu được tổng hợp từ nhật ký defect tracking thực tế của hệ thống SmartRecruit cuối Sprint 4, đã loại bỏ các báo cáo trùng lặp (duplicate); mỗi báo cáo thuộc đúng một nhóm lỗi kỹ thuật chính. Đơn vị đo là số báo cáo lỗi (Defect Count), không phải số người dùng bị ảnh hưởng hay số lần chạy test."

# Table 0: Replace with SmartRecruit defects
t0 = doc.tables[0]
defect_data = [
    ("L01", "Lỗi trích xuất & phân tích JSON từ Gemini AI trả về sai schema hoặc rỗng (AI / Backend)", "28"),
    ("L02", "Lỗi vòng lặp xác thực phiên & mất đồng bộ token trên React 19 / TanStack Query (Frontend / Auth)", "22"),
    ("L03", "Lỗi tải lên CV định dạng PDF/DOCX dung lượng lớn bị timeout / ngắt kết nối S3 (Frontend / Storage)", "16"),
    ("L04", "Lỗi tính toán sai điểm đối chiếu kỹ năng (Match Scoring) do không chuẩn hóa synonym (Backend / DB)", "14"),
    ("L06", "Lỗi Virtual Threads Java 21 bị ghim (Thread Pinning) khi gọi thư viện I/O đồng bộ (Java 21 / Backend)", "8"),
    ("L07", "Lỗi Flyway migration không tương thích cú pháp và collation trên MySQL 8.4 LTS (Database / Migration)", "5"),
    ("L08", "Lỗi vỡ layout mobile và không nhận biến CSS theme Tailwind 4 trên React 19 (Frontend / CSS)", "4"),
    ("L05", "Lỗi rò rỉ phân quyền IDOR: Người dùng xem được hồ sơ CV của ứng viên khác (Security / RBAC)", "3")
]

for idx, (code, name, count) in enumerate(defect_data, start=1):
    row = t0.rows[idx]
    row.cells[0].text = code
    row.cells[1].text = name
    row.cells[2].text = count
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = docx.shared.Pt(9.5)

# Row 9: Total
t0.rows[9].cells[0].text = ""
t0.rows[9].cells[1].text = "Tổng số báo cáo sự cố ghi nhận cuối Sprint 4"
t0.rows[9].cells[2].text = "100"
for c in t0.rows[9].cells:
    for p in c.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = docx.shared.Pt(9.5)
            r.font.bold = True

doc.paragraphs[9].text = "Thông tin bổ sung: L05 đã được đội ngũ Security xác nhận là lỗi rò rỉ phân quyền IDOR (Insecure Direct Object Reference) cực kỳ nghiêm trọng (Critical Vulnerability), vi phạm nguyên tắc Least Privilege và luật bảo vệ dữ liệu cá nhân (PII). Cần cân nhắc thông tin này khi đề xuất thứ tự xử lý, tuyệt đối không được trì hoãn chỉ vì số lượng báo cáo xuất hiện ít."
doc.paragraphs[10].text = "Chuẩn bị: Một máy tính có Excel cho mỗi nhóm. Dùng mã lỗi kèm tên nhóm lỗi thực tế của SmartRecruit để biểu đồ và kết luận dễ đối chiếu."

# 3 Phân công và quản lý thời gian
t1 = doc.tables[1]
members = [
    ("SV1", "Trần Thanh Hiệp (MSSV: 2380600636) – PM / Tech Lead", "Điều phối nhóm, tổng hợp phân tích phản biện và quyết định thứ tự ưu tiên xử lý lỗi"),
    ("SV2", "Thành viên 2 (MSSV: 2380600102) – BA / PO", "Thu thập dữ liệu lỗi Jira/GitHub, nhập bảng tính Excel và kiểm tra tính toàn vẹn"),
    ("SV3", "Thành viên 3 (MSSV: 2380600203) – Frontend UI-UX", "Thiết lập biểu đồ kết hợp Combo Chart (Pareto) trên Microsoft Excel"),
    ("SV4", "Thành viên 4 (MSSV: 2380600304) – QA Tester", "Kiểm tra công thức toán học tỷ lệ lũy kế, xác minh ngưỡng 80% và phân tích 80/20"),
    ("SV5", "Thành viên 5 (MSSV: 2380600405) – Backend Architect", "Soạn thảo kế hoạch hành động 3 công việc cải thiện chất lượng chi tiết cho Sprint 5")
]
for idx, (m_id, m_name, m_task) in enumerate(members, start=1):
    row = t1.rows[idx]
    row.cells[0].text = m_id
    row.cells[1].text = m_name
    row.cells[2].text = m_task
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = docx.shared.Pt(9.5)

# 4 Hướng dẫn thực hiện trên Excel
doc.paragraphs[17].text = "Bước 1. Tạo sheet Phan_tich với các cột A: Nhóm lỗi; B: Số báo cáo; C: Tỷ lệ; D: Tỷ lệ lũy kế; E: Ngưỡng 80%. Nhập dữ liệu 8 nhóm lỗi của SmartRecruit vào hàng 2–9. Sắp xếp cả bảng theo cột B giảm dần (từ 28 đến 3); không sắp xếp riêng một cột."
doc.paragraphs[18].text = "Bước 2. Nhập công thức dưới đây tại hàng 2 rồi sao chép đến hàng 9. Định dạng các cột C, D, E là phần trăm (0.0%)."
doc.paragraphs[21].text = "Bước 4. Đặt tên biểu đồ: “Biểu đồ Pareto phân tích lỗi hệ thống SmartRecruit cuối Sprint 4”. Hiển thị tên nhóm lỗi và chú giải. Kiểm tra tổng số báo cáo bằng 100, tổng tỷ lệ bằng 100% và lũy kế cuối cùng bằng 100%."

# 5 Phiếu trả lời của nhóm
doc.paragraphs[27].text = "Hai nhóm lỗi nhiều nhất là gì và chiếm bao nhiêu phần trăm tổng số báo cáo?"
doc.paragraphs[28].text = "Trả lời: Dựa trên bảng số liệu đã sắp xếp giảm dần, hai nhóm lỗi xuất hiện nhiều nhất trong hệ thống SmartRecruit cuối Sprint 4 là:"
doc.paragraphs[29].text = "1. L01 – Lỗi trích xuất & phân tích JSON từ Gemini AI trả về sai schema hoặc rỗng: 28 báo cáo (chiếm 28.0%).\n2. L02 – Lỗi vòng lặp xác thực phiên & mất đồng bộ token trên React 19 / TanStack Query: 22 báo cáo (chiếm 22.0%).\n-> Tổng cộng hai nhóm lỗi này chiếm 50/100 báo cáo, tương đương đúng 50.0% tổng số lượng lỗi của toàn bộ hệ thống."

doc.paragraphs[31].text = "Cần chọn tối thiểu bao nhiêu nhóm lỗi đầu tiên để lũy kế đạt ít nhất 80%? Liệt kê mã lỗi và tỷ lệ lũy kế tại điểm đó."
doc.paragraphs[32].text = "Trả lời: Cần chọn tối thiểu 4 nhóm lỗi đầu tiên để tỷ lệ tích lũy đạt ngưỡng tối thiểu 80%. Danh sách các mã lỗi và tỷ lệ lũy kế tương ứng bao gồm:"
doc.paragraphs[33].text = "- Nhóm 1 (L01): 28 báo cáo -> Tỷ lệ: 28.0% | Tích lũy: 28.0%\n- Nhóm 2 (L02): 22 báo cáo -> Tỷ lệ: 22.0% | Tích lũy: 50.0%\n- Nhóm 3 (L03 - Lỗi upload CV PDF/DOCX dung lượng lớn bị timeout S3): 16 báo cáo -> Tỷ lệ: 16.0% | Tích lũy: 66.0%\n- Nhóm 4 (L04 - Lỗi tính toán sai điểm đối chiếu kỹ năng do thiếu synonym): 14 báo cáo -> Tỷ lệ: 14.0% | Tích lũy: 80.0%\n-> Tỷ lệ lũy kế tại điểm chọn 4 nhóm lỗi đầu tiên đạt chính xác 80.0% tổng số lỗi phát sinh."

doc.paragraphs[35].text = "Dữ liệu này có đúng 20% nhóm lỗi tạo ra 80% số báo cáo không? Tính tỷ lệ số nhóm được chọn trên tổng số nhóm và giải thích."
doc.paragraphs[36].text = "Trả lời: Dữ liệu thực tế của SmartRecruit KHÔNG tuân theo tỷ lệ toán học cứng nhắc 20–80. Tỷ lệ số nhóm lỗi được chọn: 4 nhóm được chọn trên tổng số 8 nhóm lỗi = 4 / 8 = 50.0% (không phải 20%)."
doc.paragraphs[37].text = "Giải thích bản chất quản trị: Trong kỹ nghệ phần mềm và quản lý chất lượng (PMBOK Quality Management), nguyên lý Pareto không phải là một định luật toán học bất biến mà là một nguyên tắc chỉ dẫn thực nghiệm (Heuristic Rule) về 'Số ít cốt yếu và Số nhiều thứ yếu' (Vital Few vs. Trivial Many). Việc chỉ cần xử lý 4 nhóm lỗi (chiếm một nửa danh mục) đã giúp triệt tiêu tới 80% khối lượng sự cố cho thấy nguồn lực sửa lỗi trong Sprint 5 vẫn mang lại hiệu suất hoàn vốn (ROI) cực kỳ cao."

doc.paragraphs[39].text = "L05 có ít báo cáo (chỉ 3 báo cáo). Có nên trì hoãn xử lý vì nằm cuối biểu đồ không? Giải thích dựa trên tác động."
doc.paragraphs[40].text = "Trả lời: TUYỆT ĐỐI KHÔNG ĐƯỢC TRÌ HOÃN XỬ LÝ L05! Biểu đồ Pareto chỉ phản ánh khía cạnh Tần suất xuất hiện (Frequency), trong khi việc ra quyết định xử lý sự cố trong kỹ nghệ phần mềm bắt buộc phải dựa trên Ma trận rủi ro: Mức độ ưu tiên = Tần suất (Probability) × Mức độ nghiêm trọng (Severity / Impact)."
doc.paragraphs[41].text = "Đánh giá tác động của L05: L05 là lỗi rò rỉ phân quyền truy cập hồ sơ ứng viên (IDOR - Insecure Direct Object Reference), cho phép người dùng xem trộm CV của ứng viên khác. Đây là lỗi bảo mật cấp độ Critical (Blocker), vi phạm nghiêm trọng nguyên tắc Least Privilege, phá vỡ cam kết bảo vệ dữ liệu định danh cá nhân (PII) và có thể dẫn tới hậu quả pháp lý nghiêm trọng ngay cả khi chỉ xảy ra một lần duy nhất. Do đó, L05 bắt buộc phải được xếp độ ưu tiên cao nhất (P0 - Blocker) để khắc phục ngay tại đầu Sprint 5 song song với các lỗi có tần suất cao."

doc.paragraphs[43].text = "Lỗi trích xuất và phân tích JSON từ Gemini AI trả về sai schema hoặc rỗng (L01) đã phải nguyên nhân gốc chưa? Nêu ít nhất hai loại bằng chứng cần thu thập để điều tra."
doc.paragraphs[44].text = "Trả lời: L01 ('Gemini trả về sai schema hoặc rỗng') CHƯA PHẢI LÀ NGUYÊN NHÂN GỐC (Root Cause). Đây chỉ là hiện tượng/triệu chứng bề mặt (Symptom). Nguyên nhân gốc thực sự có thể nằm ở các khâu:\n1. Parser PDFBox trích xuất văn bản từ CV dạng PDF bị lỗi font encoding khiến chuỗi text gửi vào AI bị rác hoặc mất chữ.\n2. Prompt v1 chưa thiết lập chế độ Strict JSON Mode hoặc văn bản CV quá dài vượt ngưỡng token dẫn tới response bị cắt ngang.\n3. Phía Gemini 2.5 Flash đôi khi bọc output trong markdown code blocks khiến bộ deserialize Jackson của Spring Boot bị vấp ngoại lệ JsonParseException."
doc.paragraphs[45].text = "Hai loại bằng chứng cụ thể cần thu thập để điều tra nguyên nhân gốc:\n- Bằng chứng 1 (Raw Network & AI Payload Dump): Trích xuất toàn bộ chuỗi JSON raw request và raw response trao đổi giữa Spring Boot và Gemini API thông qua correlation ID / traceId được ghi nhận trên CloudWatch Logs.\n- Bằng chứng 2 (Extracted Plaintext Artifact): Lưu lại file text trung gian sau khi qua thư viện Apache PDFBox để đối chiếu xem văn bản đầu vào gửi cho AI có giữ nguyên vẹn nội dung các phần Skills và Experience hay không."

# 6 Kế hoạch cải thiện chất lượng cho Sprint 5
# Công việc 1
doc.paragraphs[51].text = "Tên công việc và mã lỗi liên quan: Cứng hóa phân quyền xác thực đa tầng và vá lỗ hổng IDOR trên API Hồ sơ ứng tuyển (Khắc phục lỗi L05)"
doc.paragraphs[52].text = "Lý do ưu tiên: Lỗ hổng bảo mật cấp độ Critical (Blocker), đe dọa vi phạm dữ liệu cá nhân PII của ứng viên và vi phạm nguyên tắc Admin/Recruiter Least Privilege. Cần xử lý triệt để ngay đầu Sprint 5 trước khi phát triển các tính năng đánh giá."
doc.paragraphs[53].text = "Người chủ trì: Trần Thanh Hiệp (Backend Architect / Tech Lead)    Người kiểm tra: SV4 (QA Tester) + Security Reviewer"
doc.paragraphs[54].text = "Điều kiện nghiệm thu và minh chứng:"
doc.paragraphs[55].text = "- 100% các endpoint /api/v1/applications/{id} và /api/v1/applications/{id}/cv được bảo vệ bởi bộ đánh giá phân quyền máy chủ ApplicationOwnershipEvaluator kết hợp @PreAuthorize.\n- Candidate chỉ truy cập được đúng application của mình; Recruiter chỉ xem được application thuộc Job mình sở hữu; Admin không thể download raw CV.\n- Minh chứng kiểm chứng: Bổ sung 15 automated negative authorization tests trong CI pipeline (đều pass 100%) và kết quả quét bảo mật OWASP ZAP không còn cảnh báo IDOR."

# Công việc 2
doc.paragraphs[57].text = "Tên công việc và mã lỗi liên quan: Cải tiến Adapter Gemini 2.5 Flash, hoàn thiện cơ chế sanitize markdown và Strict Schema Validation (Khắc phục lỗi L01)"
doc.paragraphs[58].text = "Lý do ưu tiên: Nhóm lỗi chiếm tần suất cao nhất (28% tổng số lỗi phát sinh), làm tê liệt luồng xử lý tự động trích xuất CV – tính năng giá trị cốt lõi nhất của Web MVP SmartRecruit."
doc.paragraphs[59].text = "Người chủ trì: SV5 (Backend / AI Integration Engineer)    Người kiểm tra: SV2 (BA / PO) + Trần Thanh Hiệp (Tech Lead)"
doc.paragraphs[60].text = "Điều kiện nghiệm thu và minh chứng:"
doc.paragraphs[61].text = "- Nâng cấp GeminiAdapter: tự động bóc tách các khối markdown code fence trước khi chuyển qua bộ parse Jackson.\n- Cấu hình Jackson ObjectMapper với DeserializationFeature.FAIL_ON_UNKNOWN_PROPERTIES = false và bắt chặt các trường bắt buộc; nếu thiếu trường cốt lõi, tự động gán cờ needsHumanReview = true thay vì văng Exception.\n- Minh chứng kiểm chứng: Tỷ lệ trích xuất thành công và parse hợp lệ đạt >= 95% trên bộ dữ liệu kiểm thử chuẩn 20+ CV mẫu, có file log ghi nhận metrics thời gian xử lý và token."

# Công việc 3
doc.paragraphs[63].text = "Tên công việc và mã lỗi liên quan: Xử lý Race Condition khi Refresh Token và đồng bộ trạng thái Session trên React 19 (Khắc phục lỗi L02)"
doc.paragraphs[64].text = "Lý do ưu tiên: Nhóm lỗi chiếm tần suất cao thứ hai (22% tổng số lỗi), trực tiếp phá hỏng trải nghiệm người dùng, gây văng phiên làm việc và lỗi 401 lặp vô hạn khi Access Token 15 phút hết hạn."
doc.paragraphs[65].text = "Người chủ trì: SV3 (Frontend UI-UX Engineer)    Người kiểm tra: Trần Thanh Hiệp (Tech Lead / Backend)"
doc.paragraphs[66].text = "Điều kiện nghiệm thu và minh chứng:"
doc.paragraphs[67].text = "- Triển khai Axios Interceptor với cơ chế hàng đợi Promise (Refresh Token Queue): Khi có nhiều request đồng thời gặp lỗi 401, chỉ 1 request duy nhất gọi endpoint /auth/refresh, các request khác tạm dừng chờ token mới rồi tự động thử lại (retry).\n- Đồng bộ cache TanStack Query v5 và dọn dẹp sạch sẽ session state khi người dùng thực hiện Logout hoặc bị thu hồi quyền.\n- Minh chứng kiểm chứng: Kịch bản kiểm thử tự động với Playwright mô phỏng 5 API calls đồng thời khi token hết hạn hoạt động trơn tru; người dùng thao tác liên tục trong 120 phút không bị văng ra trang Login."

# 7 Nộp bài và tiêu chí chấm điểm
doc.paragraphs[69].text = "Nộp Nhom03_Lab_Pareto.xlsx gồm sheet Phan_tich có bảng, công thức và biểu đồ; sheet Ket_luan có 6 câu trả lời, 3 công việc và đóng góp của đủ 5 thành viên. Nếu điền câu trả lời trong Word thì nộp thêm file Word này và ghi rõ trong sheet Ket_luan."

# Save to target files
doc.save(target_pareto)
print(f"SUCCESS: Saved to {target_pareto}")

doc.save(target_artifact)
print(f"SUCCESS: Saved to {target_artifact}")
