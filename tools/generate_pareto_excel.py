import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference

wb = openpyxl.Workbook()

# Sheet 1: Phan_tich
ws1 = wb.active
ws1.title = "Phan_tich"

# Sheet Title
ws1.merge_cells("A1:E1")
title_cell = ws1["A1"]
title_cell.value = "BẢNG PHÂN TÍCH LỖI HỆ THỐNG SMARTRECRUIT CUỐI SPRINT 4 (BIỂU ĐỒ PARETO)"
title_cell.font = Font(name="Arial", size=14, bold=True, color="1F4E79")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[1].height = 30

# Subtitle
ws1.merge_cells("A2:E2")
sub_cell = ws1["A2"]
sub_cell.value = "Dự án: Nền tảng tuyển dụng thông minh SmartRecruit | Thực hiện: [Nhóm 03] (Trần Thanh Hiệp - PM/Lead)"
sub_cell.font = Font(name="Arial", size=10, italic=True, color="595959")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[2].height = 20

# Headers
headers = [
    "Mã & Nhóm lỗi công nghệ thực tế (SmartRecruit MVP)",
    "Số báo cáo (Count)",
    "Tỷ lệ % (Percentage)",
    "Tỷ lệ lũy kế % (Cumulative %)",
    "Ngưỡng 80% (Threshold)"
]

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws1.row_dimensions[4].height = 26
for col_num, header in enumerate(headers, 1):
    cell = ws1.cell(row=4, column=col_num)
    cell.value = header
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

# Data rows (Sorted descending)
data = [
    ("L01 – Lỗi trích xuất & phân tích JSON từ Gemini AI trả về sai schema hoặc rỗng (AI/Backend)", 28),
    ("L02 – Lỗi vòng lặp xác thực phiên & mất đồng bộ token trên React 19 / TanStack Query (Frontend/Auth)", 22),
    ("L03 – Lỗi tải lên CV định dạng PDF/DOCX dung lượng lớn bị timeout / ngắt kết nối S3 (Frontend/Storage)", 16),
    ("L04 – Lỗi tính toán sai điểm đối chiếu kỹ năng (Match Scoring) do không chuẩn hóa synonym (Backend/DB)", 14),
    ("L06 – Lỗi Virtual Threads Java 21 bị ghim (Thread Pinning) khi gọi thư viện I/O đồng bộ (Java 21/Backend)", 8),
    ("L07 – Lỗi Flyway migration không tương thích cú pháp và collation trên MySQL 8.4 LTS (Database/Migration)", 5),
    ("L08 – Lỗi vỡ layout mobile và không nhận biến CSS theme Tailwind 4 trên React 19 (Frontend/CSS)", 4),
    ("L05 – Lỗi rò rỉ phân quyền IDOR: Người dùng xem được hồ sơ CV của ứng viên khác (Security/RBAC)", 3)
]

zebra_fill = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")
regular_font = Font(name="Arial", size=10)
vital_few_font = Font(name="Arial", size=10, bold=True, color="1F4E79")
security_alert_font = Font(name="Arial", size=10, bold=True, color="C00000")

for idx, (label, count) in enumerate(data, start=5):
    ws1.row_dimensions[idx].height = 22
    row_fill = zebra_fill if idx % 2 == 0 else PatternFill(fill_type=None)
    
    # Col A: Label
    c_a = ws1.cell(row=idx, column=1, value=label)
    c_a.alignment = Alignment(horizontal="left", vertical="center")
    c_a.border = thin_border
    c_a.fill = row_fill
    if "L05" in label:
        c_a.font = security_alert_font
    elif idx <= 8:
        c_a.font = vital_few_font
    else:
        c_a.font = regular_font
        
    # Col B: Count
    c_b = ws1.cell(row=idx, column=2, value=count)
    c_b.alignment = Alignment(horizontal="center", vertical="center")
    c_b.border = thin_border
    c_b.fill = row_fill
    c_b.font = regular_font
    
    # Col C: Formula Percentage =B2/SUM($B$2:$B$9) -> here row is idx, data is 5 to 12
    c_c = ws1.cell(row=idx, column=3, value=f"=B{idx}/SUM($B$5:$B$12)")
    c_c.number_format = "0.0%"
    c_c.alignment = Alignment(horizontal="right", vertical="center")
    c_c.border = thin_border
    c_c.fill = row_fill
    c_c.font = regular_font
    
    # Col D: Cumulative Percentage =SUM($B$5:B{idx})/SUM($B$5:$B$12)
    c_d = ws1.cell(row=idx, column=4, value=f"=SUM($B$5:B{idx})/SUM($B$5:$B$12)")
    c_d.number_format = "0.0%"
    c_d.alignment = Alignment(horizontal="right", vertical="center")
    c_d.border = thin_border
    c_d.fill = row_fill
    c_d.font = vital_few_font if idx <= 8 else regular_font
    
    # Col E: Threshold =80%
    c_e = ws1.cell(row=idx, column=5, value=0.80)
    c_e.number_format = "0.0%"
    c_e.alignment = Alignment(horizontal="right", vertical="center")
    c_e.border = thin_border
    c_e.fill = row_fill
    c_e.font = Font(name="Arial", size=10, italic=True, color="595959")

# Total Row at row 13
ws1.row_dimensions[13].height = 24
total_font = Font(name="Arial", size=11, bold=True, color="000000")
total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

c_tot_a = ws1.cell(row=13, column=1, value="TỔNG SỐ LƯỢNG BÁO CÁO LỖI (SỰ CỐ)")
c_tot_a.font = total_font
c_tot_a.fill = total_fill
c_tot_a.alignment = Alignment(horizontal="left", vertical="center")
c_tot_a.border = thin_border

c_tot_b = ws1.cell(row=13, column=2, value="=SUM(B5:B12)")
c_tot_b.font = total_font
c_tot_b.fill = total_fill
c_tot_b.alignment = Alignment(horizontal="center", vertical="center")
c_tot_b.border = thin_border

c_tot_c = ws1.cell(row=13, column=3, value="=SUM(C5:C12)")
c_tot_c.font = total_font
c_tot_c.fill = total_fill
c_tot_c.alignment = Alignment(horizontal="right", vertical="center")
c_tot_c.number_format = "0.0%"
c_tot_c.border = thin_border

c_tot_d = ws1.cell(row=13, column=4, value="100.0%")
c_tot_d.font = total_font
c_tot_d.fill = total_fill
c_tot_d.alignment = Alignment(horizontal="right", vertical="center")
c_tot_d.border = thin_border

c_tot_e = ws1.cell(row=13, column=5, value="-")
c_tot_e.font = total_font
c_tot_e.fill = total_fill
c_tot_e.alignment = Alignment(horizontal="center", vertical="center")
c_tot_e.border = thin_border

# Column widths
ws1.column_dimensions['A'].width = 80
ws1.column_dimensions['B'].width = 18
ws1.column_dimensions['C'].width = 20
ws1.column_dimensions['D'].width = 24
ws1.column_dimensions['E'].width = 20

# Create Pareto Chart
chart_col = BarChart()
chart_col.title = "Biểu đồ Pareto phân tích lỗi hệ thống SmartRecruit cuối Sprint 4"
chart_col.style = 10
chart_col.y_axis.title = "Số lượng báo cáo lỗi (Defect Count)"
chart_col.x_axis.title = "Nhóm lỗi công nghệ"
chart_col.width = 22
chart_col.height = 12

data_ref = Reference(ws1, min_col=2, min_row=4, max_row=12)
cats_ref = Reference(ws1, min_col=1, min_row=5, max_row=12)
chart_col.add_data(data_ref, titles_from_data=True)
chart_col.set_categories(cats_ref)

chart_line = LineChart()
line_data_ref = Reference(ws1, min_col=4, max_col=5, min_row=4, max_row=12)
chart_line.add_data(line_data_ref, titles_from_data=True)
chart_line.y_axis.title = "Tỷ lệ tích lũy / Ngưỡng (%)"
chart_line.y_axis.axId = 200
chart_line.y_axis.number_format = '0%'
chart_line.y_axis.majorGridlines = None

chart_col += chart_line
ws1.add_chart(chart_col, "A16")

# Sheet 2: Ket_luan
ws2 = wb.create_sheet(title="Ket_luan")

ws2.merge_cells("A1:D1")
ws2["A1"].value = "PHIẾU TRẢ LỜI CÂU HỎI VÀ KẾ HOẠCH CẢI THIỆN CHẤT LƯỢNG SPRINT 5"
ws2["A1"].font = Font(name="Arial", size=14, bold=True, color="1F4E79")
ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 30

ws2.merge_cells("A2:D2")
ws2["A2"].value = "Đơn vị thực hiện: [Nhóm 03] - Quản lý dự án CNTT | Đề tài: Nền tảng tuyển dụng thông minh SmartRecruit"
ws2["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[2].height = 20

# Q&A content
qa_rows = [
    ("Câu hỏi phân tích", "Nội dung giải trình và lập luận của nhóm [Nhóm 03]"),
    ("Câu 1: Hai nhóm lỗi nhiều nhất là gì và chiếm bao nhiêu %?",
     "1. L01 – Lỗi trích xuất & phân tích JSON từ Gemini AI trả về sai schema hoặc rỗng: 28 báo cáo (28.0%).\n"
     "2. L02 – Lỗi vòng lặp xác thực phiên & mất đồng bộ token trên React 19 / TanStack Query: 22 báo cáo (22.0%).\n"
     "-> Tổng cộng 2 nhóm lỗi này chiếm 50/100 báo cáo, tương đương đúng 50.0% tổng số sự cố ghi nhận."),
    ("Câu 2: Cần chọn tối thiểu bao nhiêu nhóm lỗi để lũy kế đạt ít nhất 80%?",
     "Cần chọn tối thiểu 4 nhóm lỗi đầu tiên để đạt đúng ngưỡng tích lũy 80.0%:\n"
     "- L01: 28% | Tích lũy: 28.0%\n"
     "- L02: 22% | Tích lũy: 50.0%\n"
     "- L03 (Upload CV timeout S3): 16% | Tích lũy: 66.0%\n"
     "- L04 (Sai điểm Match Scoring do thiếu synonym): 14% | Tích lũy: 80.0%\n"
     "-> Tại mốc 4 nhóm lỗi này, lũy kế đạt chính xác 80.0%."),
    ("Câu 3: Dữ liệu này có đúng 20% nhóm lỗi tạo ra 80% số báo cáo không? Giải thích.",
     "Tỷ lệ số nhóm được chọn là 4 / 8 nhóm = 50.0% (không phải 20%).\n"
     "Giải thích: Trong kỹ nghệ phần mềm và quản lý chất lượng PMBOK, nguyên lý Pareto là quy tắc thực nghiệm Heuristic (Số ít cốt yếu và Số nhiều thứ yếu - Vital Few vs. Trivial Many), không phải quy luật số học bất biến 20-80. Việc giải quyết 4 nhóm lỗi (1/2 danh mục) triệt tiêu được 80% sự cố toàn hệ thống vẫn đem lại tỷ suất sinh lợi đầu tư (ROI) rất lớn cho Sprint 5."),
    ("Câu 4: L05 có ít báo cáo (3 báo cáo). Có nên trì hoãn xử lý không?",
     "TUYỆT ĐỐI KHÔNG ĐƯỢC TRÌ HOÃN L05!\n"
     "Căn cứ: Biểu đồ Pareto chỉ đo Tần suất (Frequency). Quyết định xử lý sự cố phần mềm bắt buộc phải dựa trên Ma trận Rủi ro: Ưu tiên = Tần suất x Mức độ nghiêm trọng (Severity).\n"
     "L05 là lỗ hổng IDOR (Insecure Direct Object Reference) cực kỳ nghiêm trọng (Critical / Blocker), cho phép người dùng xem trái phép CV của ứng viên khác. Lỗi này vi phạm bảo mật dữ liệu PII và nguyên tắc Least Privilege, có thể gây hậu quả pháp lý nghiêm trọng ngay cả khi chỉ xuất hiện 1 lần. Do đó, L05 được gắn cờ P0 (Blocker) để vá ngay đầu Sprint 5."),
    ("Câu 5: L01 ('Gemini trả về sai schema hoặc rỗng') đã phải nguyên nhân gốc chưa? Nêu 2 loại bằng chứng cần điều tra.",
     "L01 chỉ là HIỆN TƯỢNG BỀ MẶT (Symptom), CHƯA PHẢI NGUYÊN NHÂN GỐC (Root Cause). Các nguyên nhân gốc có thể gồm: PDFBox parse lỗi font; prompt v1 chưa bật Strict JSON Mode; hoặc Gemini bao bọc mã markdown code block khiến Jackson parse văng lỗi.\n"
     "Hai loại bằng chứng cần thu thập:\n"
     "1. Bằng chứng 1 (Raw Payload Logs): Dump toàn bộ raw request và raw response giữa Spring Boot và Gemini API trên AWS CloudWatch có traceId.\n"
     "2. Bằng chứng 2 (Extracted Plaintext Artifact): Lưu file text trung gian sau khi qua PDFBox để kiểm tra văn bản đầu vào gửi cho AI có bị mất đoạn hay rác font không."),
    ("Câu 6: Kế hoạch 3 công việc cải thiện chất lượng Sprint 5",
     "Nhóm thiết lập Kế hoạch hành động 3 công việc trọng tâm (kết hợp Tần suất Pareto và Mức độ nghiêm trọng bảo mật):\n"
     "1. CÔNG VIỆC 1 (Khắc phục L05): Cứng hóa phân quyền xác thực và vá lỗ hổng IDOR trên API Hồ sơ ứng tuyển (Trần Thanh Hiệp chủ trì).\n"
     "2. CÔNG VIỆC 2 (Khắc phục L01): Cải tiến Gemini Adapter, sanitize markdown và Strict Schema Validation (Backend/AI Engineer chủ trì).\n"
     "3. CÔNG VIỆC 3 (Khắc phục L02): Xử lý Race condition khi Refresh Token và đồng bộ phiên làm việc trên React 19 (Frontend Engineer chủ trì).")
]

ws2.row_dimensions[4].height = 26
ws2.cell(row=4, column=1, value="Mục").font = header_font
ws2.cell(row=4, column=1).fill = header_fill
ws2.cell(row=4, column=1).alignment = Alignment(horizontal="center", vertical="center")
ws2.cell(row=4, column=1).border = thin_border

ws2.merge_cells("B4:D4")
h_b = ws2.cell(row=4, column=2, value="Nội dung giải trình câu hỏi phân tích Pareto của nhóm")
h_b.font = header_font
h_b.fill = header_fill
h_b.alignment = Alignment(horizontal="center", vertical="center")
h_b.border = thin_border

for r_idx, (q_t, q_a) in enumerate(qa_rows[1:], start=5):
    ws2.row_dimensions[r_idx].height = 80
    r_fill = zebra_fill if r_idx % 2 == 0 else PatternFill(fill_type=None)
    
    c1 = ws2.cell(row=r_idx, column=1, value=q_t)
    c1.font = Font(name="Arial", size=10, bold=True, color="1F4E79")
    c1.fill = r_fill
    c1.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c1.border = thin_border
    
    ws2.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx, end_column=4)
    c2 = ws2.cell(row=r_idx, column=2, value=q_a)
    c2.font = regular_font
    c2.fill = r_fill
    c2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c2.border = thin_border

ws2.column_dimensions['A'].width = 30
ws2.column_dimensions['B'].width = 35
ws2.column_dimensions['C'].width = 35
ws2.column_dimensions['D'].width = 35

output_excel = "/Users/ProM2/Documents/Nhom03_Lab_Pareto.xlsx"
wb.save(output_excel)
print(f"SUCCESS: Saved Excel workbook to {output_excel}")
