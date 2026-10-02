import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_document():
    template_path = '/Users/ProM2/Documents/Lab_QLDA_CNTT_Claude_Projects.docx'
    output_path = '/Users/ProM2/Documents/QLDA_Work.docx'

    doc = docx.Document(template_path)
    body = doc._body._element

    # Remove all existing elements except the trailing sectPr
    children = list(body)
    for c in children[:-1]:
        body.remove(c)

    # Helper: Set Header text
    sec = doc.sections[0]
    if sec.header.paragraphs:
        sec.header.paragraphs[0].text = "Lab QLDA CNTT – SmartRecruit Project"

    # Helpers for paragraph styling
    def add_p(text="", style='Normal', align=None, space_before=None, space_after=6, line_spacing=1.2, bold=False, italic=False, color=None, font_size=None):
        p = doc.add_paragraph(style=style)
        pf = p.paragraph_format
        if align is not None:
            pf.alignment = align
        if space_before is not None:
            pf.space_before = Pt(space_before)
        if space_after is not None:
            pf.space_after = Pt(space_after)
        if line_spacing is not None:
            pf.line_spacing = line_spacing
        
        if text:
            run = p.add_run(text)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            if font_size:
                run.font.size = Pt(font_size)
        return p

    def add_bullet(text_runs, space_after=3, line_spacing=1.2):
        p = doc.add_paragraph(style='List Paragraph')
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)
        pf.line_spacing = line_spacing
        
        # Add bullet XML
        numPr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr>')
        p._p.get_or_add_pPr().append(numPr)

        if isinstance(text_runs, str):
            text_runs = [(text_runs, False, False, None)]
        
        for item in text_runs:
            txt = item[0]
            b = item[1] if len(item) > 1 else False
            it = item[2] if len(item) > 2 else False
            col = item[3] if len(item) > 3 else None
            r = p.add_run(txt)
            if b: r.bold = True
            if it: r.italic = True
            if col: r.font.color.rgb = RGBColor.from_string(col)
        return p

    def add_h1(text):
        return add_p(text, style='Heading 1', bold=True)

    def add_h2(text):
        return add_p(text, style='Heading 2', bold=True)

    def add_h3(text):
        return add_p(text, style='Heading 3', bold=True)

    def add_callout(title, lines):
        t = doc.add_table(rows=1, cols=1)
        tblPr = t._tbl.tblPr
        tblPr.append(parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="A6A6A6"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="A6A6A6"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="A6A6A6"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="A6A6A6"/>
            </w:tblBorders>
        '''))
        tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="9360" w:type="dxa"/>'))
        
        cell = t.rows[0].cells[0]
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="A6A6A6"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="A6A6A6"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="A6A6A6"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="A6A6A6"/>
            </w:tcBorders>
        '''))
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F2F2F2"/>'))
        tcPr.append(parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="120" w:type="dxa"/><w:bottom w:w="120" w:type="dxa"/><w:left w:w="180" w:type="dxa"/><w:right w:w="180" w:type="dxa"/></w:tcMar>'))
        
        p0 = cell.paragraphs[0]
        p0.style = 'Normal'
        p0.paragraph_format.space_after = Pt(4)
        p0.paragraph_format.line_spacing = 1.15
        r0 = p0.add_run(title)
        r0.bold = True
        r0.font.color.rgb = RGBColor.from_string("1F4E79")
        
        for l in lines:
            p = cell.add_paragraph(style='Normal')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            if isinstance(l, str):
                p.add_run(l)
            elif isinstance(l, list):
                for run_item in l:
                    txt = run_item[0]
                    b = run_item[1] if len(run_item) > 1 else False
                    it = run_item[2] if len(run_item) > 2 else False
                    col = run_item[3] if len(run_item) > 3 else None
                    r = p.add_run(txt)
                    if b: r.bold = True
                    if it: r.italic = True
                    if col: r.font.color.rgb = RGBColor.from_string(col)
        add_p("", space_after=3)

    def add_table_data(col_widths, headers, data, alignments=None):
        t = doc.add_table(rows=len(data) + 1, cols=len(headers))
        tblPr = t._tbl.tblPr
        tblPr.append(parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>
            </w:tblBorders>
        '''))
        tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="9360" w:type="dxa"/>'))
        
        # Style Header Row
        hdr_row = t.rows[0]
        hdr_trPr = hdr_row._tr.get_or_add_trPr()
        hdr_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        hdr_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for c_idx, h_text in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{col_widths[c_idx]}" w:type="dxa"/>'))
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="1F4E79"/>'))
            tcPr.append(parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="120" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>'))
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (alignments and alignments[c_idx] == 'C') else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(h_text)
            r.bold = True
            r.font.color.rgb = RGBColor.from_string("FFFFFF")

        # Style Data Rows
        for r_idx, row_data in enumerate(data):
            row = t.rows[r_idx + 1]
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            # Alternating shading
            shd_fill = "F7FAFD" if (r_idx % 2 == 1) else "FFFFFF"

            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{col_widths[c_idx]}" w:type="dxa"/>'))
                if shd_fill != "FFFFFF":
                    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{shd_fill}"/>'))
                tcPr.append(parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="120" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>'))
                
                p = cell.paragraphs[0]
                if alignments:
                    if alignments[c_idx] == 'C':
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif alignments[c_idx] == 'R':
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                
                if isinstance(val, str):
                    p.add_run(val)
                elif isinstance(val, list):
                    for v_item in val:
                        txt = v_item[0]
                        b = v_item[1] if len(v_item) > 1 else False
                        it = v_item[2] if len(v_item) > 2 else False
                        col = v_item[3] if len(v_item) > 3 else None
                        r = p.add_run(txt)
                        if b: r.bold = True
                        if it: r.italic = True
                        if col: r.font.color.rgb = RGBColor.from_string(col)
        add_p("", space_after=3)

    return doc, add_p, add_bullet, add_h1, add_h2, add_h3, add_callout, add_table_data

print("Script framework compiled.")
