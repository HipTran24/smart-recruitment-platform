import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

wb = openpyxl.load_workbook('sheet.xlsx')
ws = wb['PERT_SmartRecruit']

# Clear old content
ws.delete_rows(1, ws.max_row)

font_title = Font(name='Arial', size=14, bold=True, color='000000')
font_header = Font(name='Arial', size=10, bold=True, color='FFFFFF')
font_bold = Font(name='Arial', size=10, bold=True, color='000000')
font_num = Font(name='Arial', size=10, bold=False, color='000000')
font_cp = Font(name='Arial', size=10, bold=True, color='DC2626')

fill_title = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
fill_header = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
fill_sprint = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
fill_cp = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')

thin_side = Side(style='thin', color='D9D9D9')
border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

# 1. Main Title
ws.cell(1, 2, 'SƠ ĐỒ TIẾN TRÌNH PERT / CPM TOÀN BỘ 50 NHIỆM VỤ DỰ ÁN SMARTRECRUIT (08/09/2026 - 03/11/2026)')
ws.cell(1, 2).font = font_title
ws.cell(1, 2).alignment = Alignment(horizontal='center', vertical='center')
ws.cell(1, 2).fill = fill_title
ws.merge_cells('B1:M1')
ws.row_dimensions[1].height = 35

# 2. Project summary banner
ws.cell(2, 2, 'Nguồn dữ liệu: 100% trích xuất từ Sơ đồ Gantt (50 SCRUM Tasks) & WBS SmartRecruit | Tổng thời hạn: 08 Sprint (57 ngày lịch, 28 ngày làm việc găng)')
ws.cell(2, 2).font = Font(name='Arial', size=10, italic=True, color='595959')
ws.cell(2, 2).alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('B2:M2')
ws.row_dimensions[2].height = 22

# 3. Table Headers
headers = [
    'Mã Ticket', 'Tên nhiệm vụ kỹ thuật (Sơ đồ Gantt)', 'Sprint', 'Bắt đầu', 'Kết thúc', 
    'Thời gian t (ngày)', 'Hoạt động tiên quyết', 'Khởi sớm (ES)', 'Kết sớm (EF)', 
    'Khởi muộn (LS)', 'Kết muộn (LF)', 'Độ trễ Slack', 'Đường găng?'
]

for col_idx, h in enumerate(headers, start=2):
    cell = ws.cell(4, col_idx, h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_all
ws.row_dimensions[4].height = 30

# Load CPM results from build_real_pert
import build_real_pert as brp

sprint_map = {
    11: 'Sprint 1', 1: 'Sprint 1', 5: 'Sprint 1', 12: 'Sprint 1', 2: 'Sprint 1', 6: 'Sprint 1', 8: 'Sprint 1', 7: 'Sprint 1', 13: 'Sprint 1', 9: 'Sprint 1', 10: 'Sprint 1', 14: 'Sprint 1', 15: 'Sprint 1',
    16: 'Sprint 2', 17: 'Sprint 2', 18: 'Sprint 2', 19: 'Sprint 2', 20: 'Sprint 2',
    3: 'Sprint 3', 21: 'Sprint 3', 22: 'Sprint 3', 23: 'Sprint 3', 24: 'Sprint 3', 25: 'Sprint 3',
    4: 'Sprint 4', 26: 'Sprint 4', 27: 'Sprint 4', 28: 'Sprint 4', 29: 'Sprint 4', 30: 'Sprint 4',
    31: 'Sprint 5', 32: 'Sprint 5', 33: 'Sprint 5', 34: 'Sprint 5', 35: 'Sprint 5',
    38: 'Sprint 6', 36: 'Sprint 6', 37: 'Sprint 6', 39: 'Sprint 6', 40: 'Sprint 6',
    41: 'Sprint 7', 42: 'Sprint 7', 44: 'Sprint 7', 43: 'Sprint 7', 45: 'Sprint 7',
    46: 'Sprint 8', 49: 'Sprint 8', 47: 'Sprint 8', 48: 'Sprint 8', 50: 'Sprint 8'
}

for idx, stt in enumerate(brp.topo_order):
    r = 5 + idx
    t = brp.tasks[stt]
    preds_str = ', '.join([f'SCRUM-{p}' for p in brp.predecessors[stt]]) if brp.predecessors[stt] else '-'
    is_cp = (brp.slack[stt] == 0)
    
    row_data = [
        t['id'],
        t['name'],
        sprint_map.get(stt, ''),
        t['start_date'],
        t['end_date'],
        t['dur'],
        preds_str,
        brp.es[stt],
        brp.ef[stt],
        brp.ls[stt],
        brp.lf[stt],
        brp.slack[stt],
        'CÓ (ĐƯỜNG GĂNG)' if is_cp else 'Không'
    ]
    
    for col_idx, val in enumerate(row_data, start=2):
        cell = ws.cell(r, col_idx, val)
        cell.border = border_all
        if is_cp:
            cell.fill = fill_cp
            cell.font = font_cp if col_idx in [2, 7, 13, 14] else font_bold
        else:
            cell.font = font_num
        
        if col_idx in [2, 4, 5, 6, 8, 9, 10, 11, 12, 13, 14]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Summary Box at the end
r_sum = 5 + len(brp.topo_order) + 2
ws.cell(r_sum, 2, 'KẾT LUẬN ĐƯỜNG GĂNG (CRITICAL PATH) DỰ ÁN SMARTRECRUIT (100% DỮ LIỆU THỰC TẾ):')
ws.cell(r_sum, 2).font = font_bold
ws.merge_cells(start_row=r_sum, start_column=2, end_row=r_sum, end_column=14)

cp_list = [f'SCRUM-{i}' for i in brp.cp_nodes]
r_sum += 1
ws.cell(r_sum, 2, ' -> '.join(cp_list))
ws.cell(r_sum, 2).font = font_cp
ws.merge_cells(start_row=r_sum, start_column=2, end_row=r_sum, end_column=14)

r_sum += 1
summary_lines = [
    f'- Tổng thời lượng làm việc trên đường găng: {brp.project_finish} ngày làm việc.',
    '- Tổng thời gian lịch trình dự án: 57 ngày (từ 08/09/2026 đến 03/11/2026 - đúng 08 Sprint).',
    '- Điểm bắt đầu: SCRUM-11 (Họp Kick-off dự án, thống nhất phạm vi, mục tiêu MVP & lập Project Charter).',
    '- Điểm kết thúc: SCRUM-50 (Họp tổng kết dự án Sprint Retrospective, nghiệm thu đồ án & đóng dự án vào 03/11/2026).'
]

for s in summary_lines:
    ws.cell(r_sum, 2, s)
    ws.cell(r_sum, 2).font = font_bold
    ws.merge_cells(start_row=r_sum, start_column=2, end_row=r_sum, end_column=14)
    r_sum += 1

# Column widths
ws.column_dimensions['A'].width = 5
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 45
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 12
ws.column_dimensions['G'].width = 14
ws.column_dimensions['H'].width = 24
ws.column_dimensions['I'].width = 12
ws.column_dimensions['J'].width = 12
ws.column_dimensions['K'].width = 12
ws.column_dimensions['L'].width = 12
ws.column_dimensions['M'].width = 14
ws.column_dimensions['N'].width = 18

wb.save('sheet.xlsx')
print('Updated sheet.xlsx with 100% REAL DATA from Sơ đồ Gantt!')
